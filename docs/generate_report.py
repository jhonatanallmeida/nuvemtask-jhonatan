"""Build the NuvemTask technical report DOCX and its architecture diagram."""

from io import BytesIO
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "relatorio-tecnico.docx"
GREEN = "28694F"
PALE = "F1F7F2"
INK = "202722"
MUTED = "5D685F"
GRID = "D9D9D9"


def font(size: int, bold: bool = False):
    file = "arialbd.ttf" if bold else "arial.ttf"
    path = Path("C:/Windows/Fonts") / file
    try:
        return ImageFont.truetype(str(path), size)
    except OSError:
        return ImageFont.load_default()


def architecture_image() -> BytesIO:
    image = Image.new("RGB", (1700, 480), "#F8FAF7")
    draw = ImageDraw.Draw(image)
    title_font = font(23, True)
    label_font = font(19)
    small_font = font(15)
    card_font = font(18, True)
    cards = [
        (50, 78, 300, "Pessoa usuária", ["Navegador", "HTTPS"]),
        (450, 78, 300, "Front-end", ["React + Vite", "Render Static Site"]),
        (850, 78, 300, "API REST", ["FastAPI + Docker", "Render Web Service"]),
        (1250, 78, 300, "Persistência", ["PostgreSQL", "Render gerenciado"]),
    ]
    for x, y, width, heading, lines in cards:
        draw.rounded_rectangle((x, y, x + width, y + 228), radius=20, fill="#FFFFFF", outline="#D4E3D7", width=3)
        draw.rounded_rectangle((x + 18, y + 18, x + 55, y + 55), radius=11, fill="#E6F1E9")
        draw.ellipse((x + 31, y + 30, x + 42, y + 41), fill="#438263")
        draw.text((x + 68, y + 21), heading, font=card_font, fill=f"#{GREEN}")
        for i, line in enumerate(lines):
            draw.text((x + 22, y + 91 + i * 34), line, font=label_font, fill=f"#{INK}")
        draw.text((x + 22, y + 187), "camada independente", font=small_font, fill="#89958B")

    arrow_y = 192
    for start, end, label in [(350, 450, "HTTPS"), (750, 850, "REST + JWT"), (1150, 1250, "SQL")]:
        draw.line((start, arrow_y, end - 10, arrow_y), fill="#6C9D7A", width=4)
        draw.polygon([(end - 10, arrow_y - 8), (end, arrow_y), (end - 10, arrow_y + 8)], fill="#6C9D7A")
        bounds = draw.textbbox((0, 0), label, font=small_font)
        label_width = bounds[2] - bounds[0]
        center = (start + end) // 2
        draw.text((center - label_width // 2, arrow_y - 30), label, font=small_font, fill="#55725C")

    ci_box = (520, 365, 1180, 455)
    draw.rounded_rectangle(ci_box, radius=16, fill="#EAF3EC", outline="#C6DDCB", width=2)
    ci_label = "GitHub Actions  ·  testes + build"
    box = draw.textbbox((0, 0), ci_label, font=title_font)
    draw.text(((1700 - (box[2] - box[0])) // 2, 392), ci_label, font=title_font, fill=f"#{GREEN}")
    # CI gates the deployment of the two web services after checks succeed.
    draw.line((680, 365, 600, 315), fill="#84A68B", width=3)
    draw.polygon([(594, 320), (600, 310), (607, 322)], fill="#84A68B")
    draw.line((1020, 365, 1000, 315), fill="#84A68B", width=3)
    draw.polygon([(994, 321), (1000, 310), (1006, 322)], fill="#84A68B")
    caption = "deploy após checks aprovados"
    box = draw.textbbox((0, 0), caption, font=small_font)
    draw.text(((1700 - (box[2] - box[0])) // 2, 337), caption, font=small_font, fill="#718276")

    buffer = BytesIO()
    image.save(buffer, format="PNG", dpi=(180, 180))
    buffer.seek(0)
    return buffer


def set_run_font(run, size=9.5, bold=False, color=INK, name="Arial"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def cell_borders(cell):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "5")
        element.set(qn("w:color"), GRID)


def cell_margins(cell, top=85, start=105, bottom=85, end=105):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{margin}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)


def add_table(document, headers, rows, widths, font_size=8.3):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for idx, (cell, value) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = Inches(widths[idx])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade_cell(cell, GREEN)
        cell_borders(cell)
        cell_margins(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(value), size=8.4, bold=True, color="FFFFFF")
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, (cell, value) in enumerate(zip(cells, row)):
            cell.width = Inches(widths[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade_cell(cell, "F4F8F4" if row_index % 2 == 0 else "FFFFFF")
            cell_borders(cell)
            cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            set_run_font(p.add_run(value), size=font_size, color=INK)
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_body(document, text, bold_lead=None):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.1
    if bold_lead and text.startswith(bold_lead):
        set_run_font(paragraph.add_run(bold_lead), size=9.4, bold=True)
        set_run_font(paragraph.add_run(text[len(bold_lead):]), size=9.4, color=MUTED)
    else:
        set_run_font(paragraph.add_run(text), size=9.4, color=MUTED)
    return paragraph


def add_heading(document, text):
    paragraph = document.add_paragraph(style="Heading 1")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(9)
    paragraph.paragraph_format.space_after = Pt(4)
    set_run_font(paragraph.add_run(text), size=13, bold=True, color="000000")
    return paragraph


def build():
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(9.4)
    normal.font.color.rgb = RGBColor.from_string(INK)
    for name in ("Title", "Heading 1", "Heading 2"):
        style = document.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.color.rgb = RGBColor(0, 0, 0)
    title_properties = document.styles["Title"]._element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)

    eyebrow = document.add_paragraph()
    eyebrow.paragraph_format.space_after = Pt(5)
    set_run_font(eyebrow.add_run("DESENVOLVIMENTO DE SOFTWARE EM NUVEM"), size=8, bold=True, color=GREEN)

    title = document.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(7)
    set_run_font(title.add_run("Relatório técnico NuvemTask"), size=23, bold=True, color="000000")

    meta = document.add_paragraph()
    meta.paragraph_format.space_after = Pt(2)
    set_run_font(meta.add_run("ADS / IA (EAD) · Unifor"), size=9.3, bold=True, color=INK)
    meta = document.add_paragraph()
    meta.paragraph_format.space_after = Pt(8)
    set_run_font(meta.add_run("Integrante: Jhonatan Almeida    |    Setembro de 2026"), size=8.5, color=MUTED)

    add_body(
        document,
        "O NuvemTask organiza projetos e tarefas com autenticação, autorização por perfil e persistência em PostgreSQL gerenciado. Este relatório apresenta a arquitetura, as tecnologias, a estratégia de CI/CD e as contribuições do desenvolvimento individual. O repositório contém a configuração de nuvem; o provisionamento público será concluído após a publicação do código e a conexão da conta Render.",
    )

    add_heading(document, "1. Visão geral do sistema")
    add_body(
        document,
        "A aplicação permite registrar usuários, criar projetos privados e manter tarefas com descrição, prazo e status (pendente, em andamento ou concluída). A pessoa usuária administra seus próprios registros; o perfil admin pode consultar todos os projetos e usuários. O sistema separa a interface web da API e do banco para facilitar manutenção e crescimento independente das camadas.",
    )

    add_heading(document, "2. Arquitetura em nuvem")
    image_paragraph = document.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.space_after = Pt(1)
    image_paragraph.add_run().add_picture(architecture_image(), width=Inches(6.83))
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(4)
    set_run_font(caption.add_run("Figura 1 — Componentes e fluxo de implantação do NuvemTask."), size=8, color=MUTED)
    add_body(
        document,
        "O React chama a API FastAPI por HTTPS, usando JSON e bearer token JWT. A API acessa o PostgreSQL por SQLAlchemy e psycopg; os dados permanecem fora do container. O serviço é stateless e pode ser replicado sem migrar sessões ou arquivos. O blueprint define uma instância inicial, e a capacidade de escala depende do plano do provedor.",
    )
    add_body(
        document,
        "O modelo de dados possui três entidades relacionadas: um usuário tem vários projetos e cada projeto tem várias tarefas. A API verifica o proprietário antes de permitir acesso ao projeto e às tarefas; a exceção de consulta ampliada é o perfil admin.",
    )

    add_heading(document, "3. Tecnologias e responsabilidades")
    add_table(
        document,
        ["Camada", "Tecnologia", "Responsabilidade"],
        [
            ["Front-end", "React 19 + Vite", "Login, painel responsivo e operações CRUD."],
            ["Back-end", "FastAPI + SQLAlchemy 2", "API REST, validação, autorização e OpenAPI."],
            ["Autenticação", "JWT HS256 + scrypt", "Token com validade de 60 min; senha derivada com salt."],
            ["Banco", "PostgreSQL gerenciado", "Persistência independente do container da API."],
            ["Container", "Docker", "Empacotamento e execução reproduzível do back-end."],
            ["CI/CD", "GitHub Actions + Render", "Testes, build e deploy após verificações aprovadas."],
            ["Observabilidade", "Logs do Python", "Acessos, status, duração, ID e erros da requisição."],
        ],
        [1.10, 1.72, 3.98],
        font_size=8.0,
    )

    add_heading(document, "4. Segurança e qualidade")
    add_body(
        document,
        "O papel não é aceito do formulário: a conta recebe `admin` somente se o e-mail corresponder a `ADMIN_EMAIL`; as demais contas são `user`. As rotas protegidas exigem token assinado com expiração. Senhas são armazenadas com scrypt e salt aleatório. O back-end valida formato e tamanho, restringe CORS às origens configuradas e responde com erro genérico em falhas internas, registrando detalhes no log.",
    )
    add_body(
        document,
        "O OpenAPI é gerado pela API e fica disponível em `/docs`. Os testes automatizados cobrem cadastro, login, perfil, CRUD de projetos e tarefas, validação, privacidade entre contas e autorização admin. Um teste de interface verifica a troca para o fluxo de cadastro; o workflow também gera o build de produção do React.",
    )

    add_heading(document, "5. Implantação e CI/CD")
    add_body(
        document,
        "O arquivo `render.yaml` descreve três recursos: serviço web Docker para a API, site estático para o front-end e PostgreSQL gerenciado. O provedor fornece a conexão do banco e gera o segredo JWT; `ADMIN_EMAIL` é configurado na criação. O front-end recebe `VITE_API_URL` no build e a origem do site é permitida pela variável CORS da API.",
    )
    add_body(
        document,
        "O workflow `.github/workflows/ci.yml` roda em pull requests e pushes para `main` ou `master`: instala dependências, executa testes de API, executa o teste de interface e compila o front-end. Com GitHub e Render conectados, `autoDeployTrigger: checksPass` libera os serviços para deploy após as verificações passarem. A URL pública e a evidência de deploy serão incluídas depois da publicação do repositório.",
    )

    add_heading(document, "6. Papéis e contribuições")
    add_table(
        document,
        ["Papel previsto na proposta", "Contribuição do integrante único"],
        [
            ["Arquiteto(a) de nuvem", "Camadas, modelo de dados e configuração do provedor."],
            ["Desenvolvedor(a) back-end", "Autenticação, API, regras, validação e logs."],
            ["Desenvolvedor(a) front-end", "Painel React, formulários e integração com a API."],
            ["Engenheiro(a) DevOps", "Docker, Compose, CI e blueprint de deploy."],
            ["Qualidade e documentação", "Testes automatizados, README, relatório e vídeo."],
        ],
        [2.10, 4.70],
        font_size=8.1,
    )

    add_heading(document, "7. Dificuldades e soluções")
    add_body(
        document,
        "A proposta prevê equipes de quatro a seis integrantes; esta entrega foi realizada individualmente por Jhonatan Almeida. Os papéis foram acumulados e registrados sem atribuir contribuições a outras pessoas. Para evitar dependência do disco efêmero do container, a persistência foi direcionada a um PostgreSQL gerenciado. Segredos e endereços ficam em variáveis de ambiente; o blueprint automatiza sua ligação sem gravar credenciais no repositório.",
    )

    ref = document.add_paragraph()
    ref.paragraph_format.space_before = Pt(2)
    ref.paragraph_format.space_after = Pt(0)
    set_run_font(ref.add_run("Referência técnica: Render, Blueprint YAML Reference — "), size=7.5, color=MUTED)
    set_run_font(ref.add_run("https://render.com/docs/blueprint-spec"), size=7.5, color=GREEN)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
